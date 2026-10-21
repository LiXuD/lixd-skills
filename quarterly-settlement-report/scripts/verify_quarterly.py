#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
季度结算报告独立复核（周期口径）

把季度报告里的每个数字，回到「报告周期」区间、从各月附件1 逐日明细重新算一遍再对撞；
同时硬校验：不得出现处罚/罚款字样。

用法：python3 verify_quarterly.py [--quarter ID] [--all]
"""
import argparse
import calendar
import json
import os
import re
import sys
from datetime import date
from statistics import mean

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(os.path.dirname(HERE), "data", "quarters.json")
ROOT = os.environ.get("MR_ROOT") or os.getcwd()
QS_ROOT = os.path.join(ROOT, "季度结算")

sys.path.insert(0, HERE)
from period_data import period_sum, period_alerts  # noqa: E402
from build_quarterly import month_path, parse_month, months_between, overlap_days  # noqa: E402


def check(q):
    start = date(*map(int, q["period_start"].split("-")))
    end = date(*map(int, q["period_end"].split("-")))
    s = period_sum(start, end)
    shifts = 0.0
    for y, m in months_between(start, end):
        rep = parse_month(month_path(y, m))
        shifts += rep["shifts"] * overlap_days(start, end, y, m) / (rep["days"] or calendar.monthrange(y, m)[1])
    shifts = round(shifts)
    alerts = period_alerts(start, end)

    path = os.path.join(QS_ROOT, q["folder"], q["out"])
    doc = docx.Document(path)
    txt = "\n".join(p.text for p in doc.paragraphs).replace(" ", "")
    cells = "\n".join(c.text for c in doc.tables[1].columns[2].cells).replace(" ", "")

    def g(pat, where):
        m = re.search(pat, where)
        return int(m.group(1).replace(",", "")) if m else None

    checks = [
        ("值机天数（周期）", (end - start).days + 1, g(r"累计值机(\d+)天", txt)),
        ("班次（按排班折算）", shifts, g(r"、(\d+)个班次，监测", txt)),
        ("交易总额", s["tot"], g(r"交易合计([\d,]+)笔", txt)),
        ("电话总数", s["call_total"], g(r"接听并记录业务来电([\d,]+)通", txt)),
        ("巡检次数（=班次）", shifts, g(r"累计巡检(\d+)次", cells)),
        ("预警起数（周期内）", len(alerts), g(r"处置预警信息(\d+)起", cells)),
        ("表格电话数", s["call_total"], g(r"共接听电话([\d,]+)个", cells)),
        ("四大系统合计", s["core"] + s["pay"] + s["wl"] + s["yl"], s["tot"]),
    ]
    for k, v in sorted(s["calls"].items(), key=lambda x: -x[1]):
        checks.append((f"电话分类·{k}", v,
                       g(re.escape(k) + r"(?:需求|问题)?([\d,]+)个（", cells)))

    cover_ok = not s["missing"]
    bad_text = [w for w in ("处罚", "罚款", "惩") if w in txt + cells]

    print(f"\n### {q['id']}  {q['period_start']}~{q['period_end']}  ({q['folder']})")
    for nm, exp, got in checks:
        print(f"  {'OK ' if exp == got else '!! '}{nm}: 周期重算={exp}  报告={got}")
    print(f"  {'OK ' if cover_ok else '!! '}附件1 周期内逐日数据完整：缺 {len(s['missing'])} 天"
          f"{'（' + str(s['missing'][:3]) + '…）' if s['missing'] else ''}")
    print(f"  {'OK ' if not bad_text else '!! '}无处罚字样：{bad_text if bad_text else '通过'}")
    print(f"  -- 交易构成 核{s['core']:,} 付{s['pay']:,} 网{s['wl']:,} 银{s['yl']:,} = {s['tot']:,}")
    print(f"  -- 批量均值 村镇 {round(mean(s['branch'])) if s['branch'] else 0} 分钟 / "
          f"XX银行 {round(mean(s['main'])) if s['main'] else 0} 分钟")
    return all(e == gv for _, e, gv in checks) and cover_ok and not bad_text


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--quarter")
    ap.add_argument("--all", action="store_true")
    a = ap.parse_args()
    cfg = json.load(open(DATA, encoding="utf-8"))
    qs = cfg["quarters"] if a.all else [q for q in cfg["quarters"] if q["id"] == a.quarter]
    if not qs:
        sys.exit("未指定或未找到季度")
    ok = all([check(q) for q in qs])
    print("\n" + "=" * 80)
    print("全部一致" if ok else "存在不一致项，必须修正后再交付")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
