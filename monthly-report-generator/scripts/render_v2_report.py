#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
月报 V2 模板渲染器
==================

以一份已成稿的 V2 格式月报为唯一排版基准（段落 XML 原样复制），
把当月数据（JSON）渲染为 V2 格式的月报 Word 文档。

设计原则
--------
1. 模板文件只读；原始月报只读；输出为新文件。
2. 正文段落直接 deepcopy 模板段落的 XML（pPr + rPr），
   因此字体、字号、行距、首行缩进、段前段后与模板逐字节一致。
3. 自动重建标题书签 _Toc243172117~129，让目录域保持可跳转。
4. 写入 <w:updateFields w:val="true"/>，Word/WPS 打开时自动刷新目录页码。

用法
----
    python render_v2_report.py 数据.json
    python render_v2_report.py 数据.json --root /path/to/月报 --out 相对路径.docx
    python render_v2_report.py 数据数组.json          # 顶层为 list 时批量渲染

JSON 结构见同目录 v2_report_schema.json。
"""
import argparse
import copy
import json
import os
import sys

from docx import Document
from docx.table import Table
from lxml import etree

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_ROOT = os.environ.get('MR_ROOT') or os.getcwd()
TEMPLATE_CANDIDATES = [
    os.environ.get('MR_TEMPLATE') or '',
    os.path.join(os.path.dirname(HERE), 'templates/月报V2模板.docx'),
]

CN_ORD = {1: '一', 2: '二', 3: '三', 4: '四', 5: '五', 6: '六', 7: '七', 8: '八'}
CN_QTY = {1: '一', 2: '两', 3: '三', 4: '四', 5: '五', 6: '六', 7: '七', 8: '八'}

# ---------- 固定正文（严格照抄 V2 模板，不得改字） ----------
FIXED_INTRO = (
    '本月围绕中心机房运维值班驻场服务，开展7×24小时值班值机、系统运行监测、'
    '机房巡检、日终批处理与备份监测、业务来电支持及异常上报跟踪工作。'
)
FIXED_STAFF = (
    '项目团队共配置9人，固定值班人数8人，分4个值班小组，'
    '采用白班、中班、夜班三班轮转方式，每班8小时，覆盖全天24小时，'
    '其余1人为现场管理及应急人员。'
)
FIXED_CHANGE = '本月无人员变动'
FIXED_SEC1 = (
    '本项目所有驻场人员均已签订保密协议并完成数据安全、操作规范、应急处置、'
    '软件合规使用相关培训，严格遵守甲方各项管理要求，'
    '本报告周期内未发生泄密、数据泄露及其他安全类重大问题。'
)
FIXED_SEC2 = (
    '本项目办公使用的电脑及所有相关软件均由甲方统一提供，不存在软件供应链安全风险，'
    '项目组人员从未擅自安装任何软件，所有软件使用均符合甲方管理要求。'
)


def q(tag):
    return W + tag


def ptext(p):
    return ''.join(t.text or '' for t in p.iter(q('t')))


def set_text(p, text):
    """保留 pPr 与首个 run 的 rPr，把段落文本整体替换。"""
    runs = p.findall(q('r'))
    if not runs:
        raise AssertionError('段落内没有 run：%r' % ptext(p))
    rpr = runs[0].find(q('rPr'))
    rpr = copy.deepcopy(rpr) if rpr is not None else None
    for r in runs:
        p.remove(r)
    new_r = etree.Element(q('r'))
    if rpr is not None:
        new_r.append(rpr)
    t = etree.SubElement(new_r, q('t'))
    t.text = text
    t.set(XML_SPACE, 'preserve')
    end = p.find(q('bookmarkEnd'))
    if end is not None:
        end.addprevious(new_r)
    else:
        p.append(new_r)


def resolve_template(explicit=None):
    for p in ([explicit] if explicit else []) + TEMPLATE_CANDIDATES:
        if p and os.path.exists(p):
            return p
    raise SystemExit('找不到 V2 模板文件，请用 --template 指定')


def build(rep, template, root, out_override=None):
    doc = Document(template)
    body = doc.element.body
    kids = list(body.iterchildren())

    # ---------- 1. 取模板原型段落 ----------
    h1_proto = h2_proto = body_proto = None
    start_idx = end_idx = None
    for i, ch in enumerate(kids):
        if ch.tag != q('p'):
            continue
        txt = ptext(ch)
        if txt == '一、本月项目概况':
            h1_proto, start_idx = ch, i
        elif txt == '（一）人员配置与值班安排':
            h2_proto = ch
        elif txt.startswith('本月围绕中心机房运维值班驻场服务'):
            body_proto = ch
        elif txt == '五、附件':
            end_idx = i
    assert h1_proto is not None, '模板中找不到标题「一、本月项目概况」'
    assert h2_proto is not None, '模板中找不到标题「（一）人员配置与值班安排」'
    assert body_proto is not None, '模板中找不到正文原型段落'
    assert end_idx is not None and end_idx > start_idx, '模板中找不到标题「五、附件」'

    anchor = kids[start_idx - 1]
    assert anchor.find('.//' + q('sectPr')) is not None, '锚点段缺少 sectPr（分节符）'

    # ---------- 2. 封面副标题 / 报告周期 ----------
    title_done = False
    for ch in kids:
        if ch.tag == q('p') and ptext(ch) == '2026年8月项目月报':
            set_text(ch, rep['title'])
            title_done = True
    assert title_done, '模板中找不到副标题「2026年8月项目月报」'

    tbl = next((ch for ch in kids if ch.tag == q('tbl')), None)
    assert tbl is not None, '模板中找不到封面信息表'
    t = Table(tbl, doc)
    assert t.cell(0, 0).text.strip().startswith('报告周期'), '封面表首行不是「报告周期」'
    set_text(t.cell(0, 1).paragraphs[0]._p, rep['period'])

    # ---------- 3. 目录条目（（四）XX月排班计划） ----------
    sdt = next((ch for ch in kids if ch.tag == q('sdt')), None)
    assert sdt is not None, '模板中找不到目录控件'
    toc_done = False
    for t_el in sdt.iter(q('t')):
        if t_el.text and '月排班计划' in t_el.text:
            t_el.text = '（四）%s排班计划' % rep['sched'][0]
            toc_done = True
    assert toc_done, '目录中找不到「（四）XX月排班计划」条目'

    # ---------- 4. 组装正文 ----------
    unit = rep.get('unit', '本月')
    total, core, pay, net, union = rep['txn']
    new_els = []
    bm = [0]

    def H(text):
        el = copy.deepcopy(h1_proto if not text.startswith('（') else h2_proto)
        bs, be = el.find(q('bookmarkStart')), el.find(q('bookmarkEnd'))
        assert bs is not None and be is not None
        name = '_Toc2431721%02d' % (17 + bm[0])
        bs.set(q('id'), str(bm[0]))
        bs.set(q('name'), name)
        be.set(q('id'), str(bm[0]))
        bm[0] += 1
        set_text(el, text)
        new_els.append(el)

    def P(text):
        el = copy.deepcopy(body_proto)
        set_text(el, text)
        new_els.append(el)

    H1 = H

    # 一、本月项目概况
    H1('一、本月项目概况')
    P(FIXED_INTRO)
    days_text = rep.get('days_text') or ('%d天' % rep['days'])
    alert_part = rep.get('alert_sentence') or (
        '%s，均已上报值班主任并协调处理' % rep['alert_summary'])
    P('%s值机覆盖%s、%d个班次、监测核心、支付、网联、银联四大系统交易合计%s笔、'
      '接听并记录业务来电%d通。%s。'
      % (unit, days_text, rep['shifts'], format(total, ','), rep['calls'], alert_part))

    # 二、人员投入与考勤
    H1('二、人员投入与考勤')
    H('（一）人员配置与值班安排')
    P(FIXED_STAFF)
    H('（二）本月人员值班考勤情况')
    P(rep['attend'])
    if rep.get('shift_para'):
        P(rep['shift_para'])
    H('（三）人员变动情况')
    P(FIXED_CHANGE)

    mon_label, year_month, sdays, mode, sw, sm, sn, savg, satt = rep['sched']
    H('（四）%s排班计划' % mon_label)
    if mode == '白夜休休':
        P('%s共%d天，值机人员排班采用白夜休休模式，确保各组工作负荷均衡。'
          '本月排班白班%d班次、夜班%d班次，值班小组平均工作天数%s天，'
          '各组班次及人员具体安排可查阅附件2《%s》。'
          % (year_month, sdays, sw, sn, savg, satt))
    else:
        P('%s共%d天，值机人员排班采用4班3倒模式，确保各组工作负荷均衡。'
          '本月排班白班%d班次、中班%d班次、夜班%d班次，值班小组平均工作天数%s天，'
          '各组班次及人员具体安排可查阅附件2《%s》。'
          % (year_month, sdays, sw, sm, sn, savg, satt))

    # 三、本月主要工作完成情况
    H1('三、本月主要工作完成情况')
    H('（一）系统交易监测情况')
    P('%s交易共计%d笔，其中核心系统共交易%d笔；支付系统共交易%d笔；'
      '网联系统共交易%d笔；银联系统共交易%d笔。每日交易详情记录可查阅附件1《%s》。'
      % (unit, total, core, pay, net, union, rep['attach1']))

    H('（二）业务来电与问题解答记录')
    P('%s共接听业务人员电话%d个，所有来电问题均已完成解答并同步反馈，问题解决率100%%。'
      '问题高发时段集中在周一至周五工作时间。' % (unit, rep['calls']))
    cat_txt = '；'.join('%s%d个，占比%s' % (n, c, p) for n, c, p in rep['cats'])
    P('按业务类型分类：%s。' % cat_txt)
    P('接听电话详细记录可查阅附件1《%s》。' % rep['attach1'])

    H('（三）数据机房值班巡检')
    alerts = rep.get('alerts') or []
    if alerts:
        P('%s机房巡检，主要存在以下%s项预警，具体情况及解决方案情况如下：'
          % (unit, CN_QTY.get(len(alerts), str(len(alerts)))))
        for i, (atitle, atime, adetail, asol) in enumerate(alerts, 1):
            P('预警%s：%s' % (CN_ORD[i], atitle))
            P('发生时间：%s' % atime)
            P('问题详情：%s' % adetail)
            P('问题解决方案：%s' % asol)
    else:
        P('%s机房巡检环境正常，未发现需上报的预警事件。' % unit)

    H('（四）日终批处理监测情况')
    c_time, c_min, w_time, w_min, note = rep['batch']
    P('%s村镇银行日终批处理监测情况如下：日终批量开始时间平均为%s，批处理平均耗时为%d分钟。'
      % (unit, c_time, c_min))
    P('%sXX银行日终批处理监测情况如下：日终批量开始时间平均为%s，批处理平均耗时为%d分钟。'
      % (unit, w_time, w_min))
    P('%s每日批处理详细时间可查阅附件1《%s》。' % (note or '', rep['attach1']))

    # 四、安全保密与数据安全
    H1('四、安全保密与数据安全')
    P(FIXED_SEC1)
    P(FIXED_SEC2)

    # 五、附件
    H1('五、附件')

    assert bm[0] == 13, '标题数量异常：%d（应为 13 个：5 个一级 + 8 个二级）' % bm[0]

    # ---------- 5. 替换正文 ----------
    for ch in kids[start_idx:end_idx + 1]:
        body.remove(ch)
    cur = anchor
    for el in new_els:
        cur.addnext(el)
        cur = el

    # ---------- 6. 打开时刷新域（目录页码） ----------
    settings = doc.settings.element
    if settings.find(q('updateFields')) is None:
        uf = etree.Element(q('updateFields'))
        uf.set(q('val'), 'true')
        compat = settings.find(q('compat'))
        if compat is not None:
            compat.addprevious(uf)
        else:
            settings.append(uf)

    out_rel = out_override or rep['out']
    out_path = out_rel if os.path.isabs(out_rel) else os.path.join(root, out_rel)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    assert os.path.abspath(out_path) != os.path.abspath(template), '不允许覆盖模板文件'
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description='按 V2 模板渲染月报 Word')
    ap.add_argument('data', help='当月数据 JSON（顶层可为对象或数组）')
    ap.add_argument('--root', default=DEFAULT_ROOT, help='月报根目录（默认 %s）' % DEFAULT_ROOT)
    ap.add_argument('--template', default=None, help='V2 模板 docx 路径')
    ap.add_argument('--out', default=None, help='输出路径（相对 --root 或绝对路径）')
    args = ap.parse_args()

    with open(args.data, encoding='utf-8') as f:
        data = json.load(f)
    reps = data if isinstance(data, list) else [data]
    template = resolve_template(args.template)

    for rep in reps:
        p = build(rep, template, args.root, args.out)
        print('OK ->', p)


if __name__ == '__main__':
    main()
