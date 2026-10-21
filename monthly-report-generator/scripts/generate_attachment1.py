#!/usr/bin/env python3
"""
生成附件1《XX月报统计数据记录表》.xlsx
- Sheet1「数据库备份详情」：从值机日志识别结果填入
- Sheet2「4大系统交易量详情」：只建表头和日期，交易量留空待用户手动填写
- Sheet3「接听电话详情」：从工作量记录 Word 解析结果填入
"""
import sys
import json
import os
import calendar
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


def avg_time_of_day(batch_data, bank, field="批量开始时间"):
    """计算某银行某时间字段的「仅时刻」平均值（丢弃日期部分）。

    ⚠️ 为什么不能直接写 =AVERAGE(B3:B32)：时间单元格存的是**含日期的完整时间值**，
    而 Excel 的 AVERAGE 是「先对序列值求和再平均」，其结果 = 日期均值的整数部分
    + 日期均值的小数部分 + 时刻均值。只有当各日日期连续、且行数为奇数时，
    日期均值的小数部分才恰好为 0、结果才等于时刻均值。否则会被污染：
      · 30 天的月份 → 日期均值小数部分 = 0.5 → 结果整体偏 12 小时（如 23:14 显示成 11:14）；
      · 有缺日/重复日（如 8 月缺 8/27、8/28 出现两次）→ 即使 31 行也会偏（8月村镇 22:43 vs 正确 21:57）。
    因此此处一律预计算「仅时刻」均值并写入静态值。
    """
    secs = []
    for rec in (batch_data or {}).values():
        info = rec.get(bank) or {}
        v = info.get(field)
        if v is None:
            continue
        if isinstance(v, str):
            parsed = None
            for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%H:%M:%S", "%H:%M"):
                try:
                    parsed = datetime.strptime(v, fmt)
                    break
                except ValueError:
                    continue
            if parsed is None:
                continue
            v = parsed
        if isinstance(v, datetime):
            secs.append(v.hour * 3600 + v.minute * 60 + v.second)
    if not secs:
        return None
    total = int(round(sum(secs) / len(secs))) % 86400
    return datetime.min.time().replace(hour=total // 3600,
                                       minute=(total % 3600) // 60,
                                       second=total % 60)


def create_attachment1(month_num, year, batch_data=None, phone_data=None, output_path=None):
    """
    生成附件1
    
    参数:
        month_num: 月份(1-12)
        year: 年份(2026)
        batch_data: 值机日志识别结果 {日期: {"村镇银行": {...}, "XX银行": {...}}}
        phone_data: 工作量记录解析结果 {日期: {柜员强制签退: n, ...}}
        output_path: 输出路径，若为None则自动生成
    """
    wb = Workbook()
    days_in_month = calendar.monthrange(year, month_num)[1]
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    bold_font = Font(bold=True)
    
    if batch_data is None:
        batch_data = {}
    if phone_data is None:
        phone_data = {}
    
    # ===== Sheet1: 数据库备份详情 =====
    ws1 = wb.active
    ws1.title = "数据库备份详情"
    
    # 村镇银行部分（行 1-33）
    # 行1-2: 表头
    ws1.merge_cells('A1:A2')
    ws1['A1'] = '日期'
    ws1['A1'].font = bold_font
    ws1['A1'].alignment = center_align
    ws1['A1'].border = thin_border
    
    ws1.merge_cells('B1:H1')
    ws1['B1'] = '村镇银行'
    ws1['B1'].font = bold_font
    ws1['B1'].alignment = center_align
    ws1['B1'].border = thin_border
    
    cz_headers = ['批量开始时间', '批量结束时间', '耗时（分钟）', '日切完成时间', 
                  'ogg备份时间', 'ogg备份时间结束时间', 'ogg备份耗时']
    for i, h in enumerate(cz_headers):
        cell = ws1.cell(row=2, column=i+2, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
    
    # 村镇银行数据行
    for day in range(1, days_in_month + 1):
        row = day + 2
        date_str = f"{month_num}月{day}号"
        ws1.cell(row=row, column=1, value=date_str).border = thin_border
        ws1.cell(row=row, column=1).alignment = center_align
        
        date_key = f"{year}{month_num:02d}{day:02d}"
        cz_info = batch_data.get(date_key, {}).get("村镇银行", {})
        
        if cz_info:
            # 批量开始时间
            val = cz_info.get("批量开始时间")
            cell = ws1.cell(row=row, column=2, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            if cz_info.get("异常"):
                cell.font = Font(color="FF0000", bold=True)
            
            # 批量结束时间
            val = cz_info.get("批量结束时间")
            cell = ws1.cell(row=row, column=3, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            if cz_info.get("异常"):
                cell.font = Font(color="FF0000", bold=True)
            
            # 耗时
            val = cz_info.get("耗时")
            cell = ws1.cell(row=row, column=4, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            if cz_info.get("异常"):
                cell.font = Font(color="FF0000", bold=True)
            
            # 日切完成时间
            val = cz_info.get("日切完成时间")
            cell = ws1.cell(row=row, column=5, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            
            # ogg备份时间
            val = cz_info.get("ogg备份时间")
            cell = ws1.cell(row=row, column=6, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            
            # ogg备份结束时间
            val = cz_info.get("ogg备份结束时间")
            cell = ws1.cell(row=row, column=7, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            
            # ogg备份耗时
            val = cz_info.get("ogg备份耗时")
            cell = ws1.cell(row=row, column=8, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            if cz_info.get("ogg异常"):
                cell.font = Font(color="FF0000", bold=True)
        else:
            for col in range(2, 9):
                ws1.cell(row=row, column=col).border = thin_border
    
    # 平均时长行
    avg_row = days_in_month + 3
    ws1.cell(row=avg_row, column=1, value='平均时长').font = bold_font
    ws1.cell(row=avg_row, column=1).border = thin_border
    # 批量开始时间平均：写「仅时刻平均」预计算值（不能用 AVERAGE，原因见 avg_time_of_day 注释）
    avg_b = avg_time_of_day(batch_data, "村镇银行")
    cell = ws1.cell(row=avg_row, column=2, value=avg_b)
    cell.border = thin_border
    cell.number_format = 'hh:mm'
    ws1.cell(row=avg_row, column=4, value=f'=AVERAGE(D3:D{days_in_month+2})').border = thin_border
    ws1.cell(row=avg_row, column=8, value=f'=AVERAGE(H3:H{days_in_month+2})').border = thin_border
    
    # 空行
    blank_row = avg_row + 1
    
    # 主行（XX银行）部分
    main_start_row = blank_row + 1  # 行35左右
    
    # 合并标题
    ws1.merge_cells(f'A{main_start_row}:A{main_start_row+1}')
    ws1.cell(row=main_start_row, column=1, value='日期').font = bold_font
    ws1.cell(row=main_start_row, column=1).alignment = center_align
    ws1.cell(row=main_start_row, column=1).border = thin_border
    
    ws1.merge_cells(f'B{main_start_row}:K{main_start_row}')
    ws1.cell(row=main_start_row, column=2, value='XX银行').font = bold_font
    ws1.cell(row=main_start_row, column=2).alignment = center_align
    ws1.cell(row=main_start_row, column=2).border = thin_border
    
    main_headers = ['批量开始时间', '批量结束时间', '耗时（分钟）', '日切完成时间',
                  'ogg2备份时间', 'ogg2备份结束时间', 'ogg2备份时间耗时',
                  'ogg4备份时间', 'ogg4备份结束时间', 'ogg4备份时间耗时']
    for i, h in enumerate(main_headers):
        cell = ws1.cell(row=main_start_row+1, column=i+2, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
    
    # 主行数据行
    for day in range(1, days_in_month + 1):
        row = main_start_row + 1 + day
        date_str = f"{month_num}月{day}号"
        ws1.cell(row=row, column=1, value=date_str).border = thin_border
        ws1.cell(row=row, column=1).alignment = center_align
        
        date_key = f"{year}{month_num:02d}{day:02d}"
        main_info = batch_data.get(date_key, {}).get("XX银行", {})
        
        if main_info:
            col_idx = 2
            for field in ["批量开始时间", "批量结束时间", "耗时", "日切完成时间",
                         "ogg2备份时间", "ogg2备份结束时间", "ogg2备份耗时",
                         "ogg4备份时间", "ogg4备份结束时间", "ogg4备份耗时"]:
                val = main_info.get(field)
                cell = ws1.cell(row=row, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = center_align
                if main_info.get("异常") and field in ["批量开始时间", "批量结束时间", "耗时"]:
                    cell.font = Font(color="FF0000", bold=True)
                col_idx += 1
        else:
            for col in range(2, 12):
                ws1.cell(row=row, column=col).border = thin_border
    
    # 主行平均时长行
    main_avg_row = main_start_row + 1 + days_in_month + 1
    ws1.cell(row=main_avg_row, column=1, value='平均时长').font = bold_font
    ws1.cell(row=main_avg_row, column=1).border = thin_border
    data_start = main_start_row + 2
    data_end = main_start_row + 1 + days_in_month
    # 批量开始时间平均：写「仅时刻平均」预计算值（不能用 AVERAGE，原因见 avg_time_of_day 注释）
    avg_b = avg_time_of_day(batch_data, "XX银行")
    cell = ws1.cell(row=main_avg_row, column=2, value=avg_b)
    cell.border = thin_border
    cell.number_format = 'hh:mm'
    ws1.cell(row=main_avg_row, column=4, value=f'=AVERAGE(D{data_start}:D{data_end})').border = thin_border
    ws1.cell(row=main_avg_row, column=8, value=f'=AVERAGE(H{data_start}:H{data_end})').border = thin_border
    ws1.cell(row=main_avg_row, column=11, value=f'=AVERAGE(K{data_start}:K{data_end})').border = thin_border
    
    # ===== Sheet2: 4大系统交易量详情 =====
    ws2 = wb.create_sheet("4大系统交易量详情")
    
    # 表头
    headers2 = ['日期', '核心系统', '支付系统', '网联系统', '银联系统']
    for i, h in enumerate(headers2):
        cell = ws2.cell(row=1, column=i+1, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
    
    # 日期行（交易量留空，用户手动填写）
    for day in range(1, days_in_month + 1):
        row = day + 1
        ws2.cell(row=row, column=1, value=f"{month_num}月{day}号").border = thin_border
        ws2.cell(row=row, column=1).alignment = center_align
        for col in range(2, 6):
            ws2.cell(row=row, column=col).border = thin_border
    
    # 汇总行
    sum_row = days_in_month + 2
    ws2.cell(row=sum_row, column=1, value='交易笔数总数').font = bold_font
    ws2.cell(row=sum_row, column=1).border = thin_border
    for col in range(2, 6):
        col_letter = get_column_letter(col)
        ws2.cell(row=sum_row, column=col, value=f'=SUM({col_letter}2:{col_letter}{days_in_month+1})').border = thin_border
    
    avg_row2 = sum_row + 1
    ws2.cell(row=avg_row2, column=1, value='交易笔数日评价数').font = bold_font
    ws2.cell(row=avg_row2, column=1).border = thin_border
    for col in range(2, 6):
        col_letter = get_column_letter(col)
        ws2.cell(row=avg_row2, column=col, value=f'=AVERAGE({col_letter}2:{col_letter}{days_in_month+1})').border = thin_border
    
    # ===== Sheet3: 接听电话详情 =====
    ws3 = wb.create_sheet("接听电话详情")
    
    # 表头
    phone_headers = ['日期', '柜员强制签退', '终端变更', '数据查询业务', '人为操作失误',
                     '设备故障异常', '系统异常故障', '银行业务问题', '机房环境',
                     '服务器设备问题', '网络设备问题', '基础设施问题']
    for i, h in enumerate(phone_headers):
        cell = ws3.cell(row=1, column=i+1, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
    
    # 数据行
    for day in range(1, days_in_month + 1):
        row = day + 1
        date_str = f"{month_num}月{day}号"
        ws3.cell(row=row, column=1, value=date_str).border = thin_border
        ws3.cell(row=row, column=1).alignment = center_align
        
        date_key = f"{year}{month_num:02d}{day:02d}"
        day_data = phone_data.get(date_key, {})
        
        for col_idx, field in enumerate(phone_headers[1:], 2):
            val = day_data.get(field, None)
            cell = ws3.cell(row=row, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = center_align
    
    # 汇总行
    phone_sum_row = days_in_month + 2
    ws3.cell(row=phone_sum_row, column=1, value='各类问题小计').font = bold_font
    ws3.cell(row=phone_sum_row, column=1).border = thin_border
    for col in range(2, 13):
        col_letter = get_column_letter(col)
        ws3.cell(row=phone_sum_row, column=col, value=f'=SUM({col_letter}2:{col_letter}{days_in_month+1})').border = thin_border
    
    phone_total_row = phone_sum_row + 1
    ws3.cell(row=phone_total_row, column=1, value='处理问题总计').font = bold_font
    ws3.cell(row=phone_total_row, column=1).border = thin_border
    ws3.cell(row=phone_total_row, column=2, value=f'=SUM(B{phone_sum_row}:L{phone_sum_row})').border = thin_border
    
    # 设置列宽
    ws1.column_dimensions['A'].width = 12
    for col in range(2, 12):
        ws1.column_dimensions[get_column_letter(col)].width = 18
    
    ws2.column_dimensions['A'].width = 18
    for col in range(2, 6):
        ws2.column_dimensions[get_column_letter(col)].width = 15
    
    ws3.column_dimensions['A'].width = 15
    for col in range(2, 13):
        ws3.column_dimensions[get_column_letter(col)].width = 15
    
    # 保存
    if output_path is None:
        output_path = f"附件1《{year}年{month_num}月报统计数据记录表》.xlsx"
    
    wb.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 3:
        print("用法: python generate_attachment1.py <月份> <年份> [batch_json] [phone_json] [output_path]")
        print("  batch_json: 值机日志识别结果的JSON文件路径（可选）")
        print("  phone_json: 工作量记录解析结果的JSON文件路径（可选）")
        sys.exit(1)
    
    month_num = int(sys.argv[1])
    year = int(sys.argv[2])
    
    batch_data = None
    phone_data = None
    output_path = None
    
    if len(sys.argv) > 3:
        with open(sys.argv[3], 'r', encoding='utf-8') as f:
            batch_data = json.load(f)
    
    if len(sys.argv) > 4:
        with open(sys.argv[4], 'r', encoding='utf-8') as f:
            phone_data = json.load(f)
    
    if len(sys.argv) > 5:
        output_path = sys.argv[5]
    
    result = create_attachment1(month_num, year, batch_data, phone_data, output_path)
    print(f"附件1已生成: {result}")


if __name__ == "__main__":
    main()
