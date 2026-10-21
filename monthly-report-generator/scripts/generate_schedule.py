#!/usr/bin/env python3
"""
生成附件2《XX月排班表》.xlsx
按 3白→3中→3夜→3休 规律从上月排班续接，项目经理按国家节假日排班

排班规则：
- 4组人员，每组12天一个完整周期：3天白班→3天中班→3天夜班→3天休息
- 任何时刻：3组在上班（白/中/夜），1组在休息
- 上完3天某个班次后，该组进入下一个班次/休息

格式规则：
- 休息日（周末+法定节假日）整行添加淡蓝色底纹（钢蓝，着色1，浅色60%，#B6C7EA）
- 工作日（含调休上班的周末）无底纹
- 项目经理按国家节假日排班（使用chinesecalendar库判断）
  - 默认周一至五上班
  - 法定节假日休息（如端午节、国庆等）
  - 调休日上班（如国庆后周六补班）

存放规则：排班表放在上个月的文件夹中
  - 7月排班表 → 保存在 26年6月/ 文件夹
  - 6月排班表 → 保存在 26年5月/ 文件夹

支持两种调用方式：
1. 指定上月排班表路径：python generate_schedule.py 7 2026 <上月排班表路径> <输出路径>
2. 自动查找上月排班表：python generate_schedule.py 7 2026
"""
import sys
import os
import glob
import calendar
from datetime import datetime, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


# 基础工作目录：优先取环境变量 MR_ROOT，其次取当前工作目录
BASE_DIR = os.environ.get("MR_ROOT") or os.getcwd()

# 4组人员的轮转顺序（按自己的班组人员替换；一组两人）
GROUPS = [
    "值班员A1、值班员A2",   # 0
    "值班员B1、值班员B2",   # 1
    "值班员C1、值班员C2",   # 2
    "值班员D1、值班员D2",   # 3
]

# 现场常班人员（项目经理），按国家节假日排班
MANAGER = "项目经理"

# 休息日底纹色：钢蓝，着色1，浅色60%（从WPS主题色accent1=#4874CB tint=0.6计算得到）
REST_DAY_FILL = PatternFill(start_color='B6C7EA', end_color='B6C7EA', fill_type='solid')

# 每组的个人周期：3白→3中→3夜→3休，共12天
CYCLE_DAYS = 12
SHIFTS_IN_CYCLE = ['白', '白', '白', '中', '中', '中', '夜', '夜', '夜', '休', '休', '休']


def prev_month(month_num, year):
    """返回上个月的月份和年份"""
    m = month_num - 1
    y = year
    if m == 0:
        m = 12
        y = year - 1
    return m, y


def find_month_dir(month_num, year):
    """定位月份工作目录（兼容26年4月和2026年4月两种格式）"""
    short_year = str(year)[2:]
    dir_path = os.path.join(BASE_DIR, f"{short_year}年{month_num}月")
    if os.path.exists(dir_path):
        return dir_path
    dir_path = os.path.join(BASE_DIR, f"{year}年{month_num}月")
    if os.path.exists(dir_path):
        return dir_path
    return None


def ensure_month_dir(month_num, year):
    """确保月份目录存在，不存在则创建，返回目录路径"""
    existing = find_month_dir(month_num, year)
    if existing:
        return existing
    short_year = str(year)[2:]
    dir_path = os.path.join(BASE_DIR, f"{short_year}年{month_num}月")
    os.makedirs(dir_path, exist_ok=True)
    print(f"已创建目录: {dir_path}")
    return dir_path


def find_prev_schedule(target_month, target_year):
    """
    自动查找上月的排班表文件
    
    存放规则：排班表放在上个月的文件夹中
    生成7月排班表 → 找6月排班表 → 在5月文件夹
    """
    prev_schedule_month = target_month - 1
    prev_schedule_year = target_year
    if prev_schedule_month == 0:
        prev_schedule_month = 12
        prev_schedule_year = target_year - 1
    
    # 上月排班表所在的文件夹 = 排班表月份再减1
    folder_month, folder_year = prev_month(prev_schedule_month, prev_schedule_year)
    
    folder_dir = find_month_dir(folder_month, folder_year)
    if not folder_dir:
        print(f"警告：未找到排班表所在目录({folder_year}年{folder_month}月)")
        return None
    
    # 精确搜索
    short_prev_year = str(prev_schedule_year)[2:]
    target_filename = f"附件2《{short_prev_year}年{prev_schedule_month}月排班表》.xlsx"
    exact_path = os.path.join(folder_dir, target_filename)
    
    if os.path.exists(exact_path):
        print(f"找到上月排班表: {exact_path}")
        return exact_path
    
    # 模糊搜索
    for pattern in ["附件2*排班表*.xlsx", "*排班表*.xlsx"]:
        matches = glob.glob(os.path.join(folder_dir, pattern))
        if matches:
            result = [m for m in matches if "附件2" in os.path.basename(m)] or matches
            print(f"找到上月排班表: {result[0]}")
            return result[0]
    
    # 备选：在 target_month-1 的文件夹里找
    alt_folder_month, alt_folder_year = prev_month(target_month, target_year)
    alt_dir = find_month_dir(alt_folder_month, alt_folder_year)
    if alt_dir:
        matches = glob.glob(os.path.join(alt_dir, "*排班表*.xlsx"))
        if matches:
            print(f"在备选目录找到排班表: {matches[0]}")
            return matches[0]
    
    print("警告：未找到上月排班表")
    return None


def get_group_index(group_name):
    """根据组员名字找到组索引"""
    first_name = group_name.split("、")[0] if group_name else ""
    for i, g in enumerate(GROUPS):
        if first_name in g:
            return i
    return -1


def get_last_day_state(prev_schedule_path):
    """
    从上月排班表中获取最后一天每个组的周期位置
    
    核心逻辑：
    - 4组同进同退，所有组在3天段中的第几天是一样的
    - 只需算白班组连续上了几天白班，就能推算所有组的位置
    
    返回: dict {group_idx: cycle_position}
    - cycle_position: 0-11, 其中 0-2=白, 3-5=中, 6-8=夜, 9-11=休
    """
    wb = load_workbook(prev_schedule_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    
    # 找最后数据行
    last_row = ws.max_row
    while last_row > 3:
        if ws.cell(row=last_row, column=1).value is not None:
            break
        last_row -= 1
    
    # 读取最后一天的白/中/夜
    last_white = ws.cell(row=last_row, column=3).value
    last_mid = ws.cell(row=last_row, column=4).value
    last_night = ws.cell(row=last_row, column=5).value
    wb.close()
    
    if not last_white:
        return None
    
    # 确定白班组在3天段的第几天（1-indexed）
    # 往回看，白班组连续几天在白班
    white_idx = get_group_index(last_white)
    
    # 重新打开读取前面的行
    wb = load_workbook(prev_schedule_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    
    days_in_shift = 1  # 至少1天
    for row in range(last_row - 1, max(3, last_row - 3), -1):
        prev_white = ws.cell(row=row, column=3).value
        if prev_white and get_group_index(prev_white) == white_idx:
            days_in_shift += 1
        else:
            break
    wb.close()
    
    # day_in_segment: 0-indexed (0,1,2)
    day_in_segment = days_in_shift - 1
    
    # 根据白班组位置推算所有组的周期位置
    # 白班组: base=0, 中班组: base=3, 夜班组: base=6, 休息组: base=9
    shift_to_base = {'白': 0, '中': 3, '夜': 6, '休': 9}
    
    mid_idx = get_group_index(last_mid) if last_mid else -1
    night_idx = get_group_index(last_night) if last_night else -1
    
    # 找休息组
    working = {white_idx, mid_idx, night_idx}
    rest_idx = -1
    for i in range(len(GROUPS)):
        if i not in working:
            rest_idx = i
            break
    
    # 构建每个组的 cycle_position
    result = {}
    result[white_idx] = shift_to_base['白'] + day_in_segment
    if mid_idx >= 0:
        result[mid_idx] = shift_to_base['中'] + day_in_segment
    if night_idx >= 0:
        result[night_idx] = shift_to_base['夜'] + day_in_segment
    if rest_idx >= 0:
        result[rest_idx] = shift_to_base['休'] + day_in_segment
    
    return result


def get_group_cycle_state(group_idx, initial_position, day_offset):
    """
    计算某个组在 day_offset 天后的班次
    
    参数:
        group_idx: 组索引
        initial_position: 该组在周期中的起始位置 (0-11)
        day_offset: 天数偏移 (0=第一天, 1=第二天, ...)
    
    返回: '白'/'中'/'夜'/'休'
    """
    pos = (initial_position + day_offset) % CYCLE_DAYS
    return SHIFTS_IN_CYCLE[pos]


def generate_schedule(month_num, year, prev_schedule_path=None, output_path=None):
    """
    生成排班表
    
    排班规则：3白→3中→3夜→3休（12天一个周期）
    每组在这个周期中轮转，4组错开3天
    
    存放规则：排班表放在上个月的文件夹中
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    days_in_month = calendar.monthrange(year, month_num)[1]
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    bold_font = Font(bold=True)
    
    # 确定每组在周期中的起始位置
    group_positions = None
    if prev_schedule_path and os.path.exists(prev_schedule_path):
        group_positions = get_last_day_state(prev_schedule_path)
        if group_positions:
            # 下一天：每组位置+1
            group_positions = {k: (v + 1) % CYCLE_DAYS for k, v in group_positions.items()}
            print(f"从上月排班续接，各组周期起始位置: {group_positions}")
        else:
            print("警告：无法解析上月排班表，使用默认起始位置")
    
    if group_positions is None:
        # 默认：组0从白班第1天开始，其他组依次错开3天
        # 组0: pos=0(白), 组1: pos=9(休), 组2: pos=6(夜), 组3: pos=3(中)
        # 这样第一天：白=0, 中=3, 夜=2, 休=1
        group_positions = {0: 0, 1: 9, 2: 6, 3: 3}
        print("使用默认起始位置")
    
    # 标题行
    ws.merge_cells('A1:F1')
    short_year = str(year)[2:]
    ws['A1'] = f'{short_year}年{month_num}月排班表'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # 表头行
    headers = ['日  期', '星期', '白班\n(9:00--17:00)', '中班\n（17:00--次日01:00）', 
               '夜班\n(次日01:00--次日09:00)', '常班\n(08:00-18:00)']
    for i, h in enumerate(headers):
        cell = ws.cell(row=2, column=i+1, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
    
    # 子表头
    sub_headers = ['', '', '公司人员', '公司人员', '公司人员', '公司人员现场项目经理']
    for i, h in enumerate(sub_headers):
        cell = ws.cell(row=3, column=i+1, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
    
    weekday_names = ['星期一', '星期二', '星期三', '星期四', '星期五', '星期六', '星期日']
    
    for day in range(1, days_in_month + 1):
        row = day + 3
        date = datetime(year, month_num, day)
        day_offset = day - 1  # 0-indexed
        
        # 日期
        cell = ws.cell(row=row, column=1, value=date)
        cell.number_format = 'YYYY-MM-DD'
        cell.border = thin_border
        cell.alignment = center_align
        
        # 星期
        weekday = weekday_names[date.weekday()]
        cell = ws.cell(row=row, column=2, value=weekday)
        cell.border = thin_border
        cell.alignment = center_align
        
        # 计算每个组今天的班次
        today_shifts = {}  # {group_idx: '白'/'中'/'夜'/'休'}
        for gidx, init_pos in group_positions.items():
            shift = get_group_cycle_state(gidx, init_pos, day_offset)
            today_shifts[gidx] = shift
        
        # 填入白班、中班、夜班
        for shift_name, col in [('白', 3), ('中', 4), ('夜', 5)]:
            assigned = None
            for gidx, shift in today_shifts.items():
                if shift == shift_name:
                    assigned = GROUPS[gidx]
                    break
            cell = ws.cell(row=row, column=col, value=assigned)
            cell.border = thin_border
            cell.alignment = center_align
        
        # 常班：现场项目经理，按国家节假日排班
        # 使用chinesecalendar库判断：工作日上班，节假日/周末休息，调休日上班
        try:
            import chinese_calendar
            is_workday = chinese_calendar.is_workday(date.date())
        except (ImportError, NotImplementedError):
            # 回退：如果库不可用或日期超出范围，按周一至五上班
            is_workday = date.weekday() < 5
        
        cell = ws.cell(row=row, column=6, value=MANAGER if is_workday else None)
        cell.border = thin_border
        cell.alignment = center_align
        
        # 休息日整行添加淡蓝色底纹（工作日无底纹）
        # 休息日 = 非工作日（含周末和法定节假日），工作日 = 含调休上班的周末
        if not is_workday:
            for col in range(1, 7):
                ws.cell(row=row, column=col).fill = REST_DAY_FILL
    
    # 列宽
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 22
    ws.column_dimensions['F'].width = 22
    
    # 确定输出路径
    if output_path is None:
        folder_month, folder_year = prev_month(month_num, year)
        month_dir = ensure_month_dir(folder_month, folder_year)
        filename = f"附件2《{short_year}年{month_num}月排班表》.xlsx"
        output_path = os.path.join(month_dir, filename)
    
    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
    
    wb.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 3:
        print("用法:")
        print("  python generate_schedule.py <月份> <年份> [上月排班表路径] [输出路径]")
        print("")
        print("存放规则：排班表放在上个月的文件夹中")
        print("  - 7月排班表 → 26年6月/附件2《26年7月排班表》.xlsx")
        print("  - 6月排班表 → 26年5月/附件2《26年6月排班表》.xlsx")
        print("")
        print("排班规则：3白→3中→3夜→3休（12天一周期），4组错开3天")
        print("")
        print("示例:")
        print("  python generate_schedule.py 7 2026")
        print("  python generate_schedule.py 7 2026 /path/to/prev.xlsx")
        print("  python generate_schedule.py 7 2026 /path/to/prev.xlsx /path/to/output.xlsx")
        sys.exit(1)
    
    month_num = int(sys.argv[1])
    year = int(sys.argv[2])
    prev_path = sys.argv[3] if len(sys.argv) > 3 else None
    output_path = sys.argv[4] if len(sys.argv) > 4 else None
    
    if prev_path is None:
        prev_path = find_prev_schedule(month_num, year)
    
    result = generate_schedule(month_num, year, prev_path, output_path)
    print(f"排班表已生成: {result}")


if __name__ == "__main__":
    main()
